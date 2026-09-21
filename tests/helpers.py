"""Builders for binary test fixtures (created in memory, so the repo stays tiny)."""

from __future__ import annotations

import io
import zipfile


def make_pdf(pages: list[str]) -> bytes:
    """A minimal, valid PDF with one text line per ``\\n`` (empty string = blank page)."""
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [{}] /Count {} >>".format(
            " ".join(f"{4 + 2 * i} 0 R" for i in range(len(pages))), len(pages)
        ),
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for index, text in enumerate(pages):
        objects.append(
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * index} 0 R >>"
        )
        body = " ".join(f"({line}) Tj 0 -16 Td" for line in text.split("\n")) if text else ""
        stream = f"BT /F1 12 Tf 72 720 Td {body} ET" if text else ""
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")

    out = b"%PDF-1.4\n"
    offsets = []
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{obj}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


def make_docx() -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("SERVICE AGREEMENT", level=1)
    document.add_paragraph("The contractor shall deliver the software by 30 June 2026.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Milestone"
    table.cell(0, 1).text = "Payment"
    table.cell(1, 0).text = "Delivery"
    table.cell(1, 1).text = "Rs. 1,00,000"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Payments"
    sheet.append(["Month", "Amount", "Status"])
    sheet.append(["April", 25000, "Paid"])
    sheet.append(["May", 25000, "Late"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_pptx() -> bytes:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Policy update"
    slide.placeholders[1].text = "Refunds are processed within 14 days."
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def make_zip(entries: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    return buffer.getvalue()


def make_wav(seconds: float = 0.5, rate: int = 16000) -> bytes:
    import struct

    samples = int(seconds * rate)
    header = b"RIFF" + struct.pack("<I", 36 + samples * 2) + b"WAVEfmt "
    header += struct.pack("<IHHIIHH", 16, 1, 1, rate, rate * 2, 2, 16)
    return header + b"data" + struct.pack("<I", samples * 2) + b"\x00\x00" * samples


PNG_BYTES = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64

"""Unit tests for service internals: Gemini wrapper helpers, export and document store."""

from __future__ import annotations

import io

import httpx
import numpy as np
import pytest
from docx import Document
from google.genai import errors as genai_errors
from google.genai import types

from app.config import Settings
from app.services.documents import StoredDocument
from app.services.export import markdown_to_docx
from app.services.llm import (
    ChatMessage,
    GeminiClient,
    LLMError,
    ModelRouter,
    _LRUCache,
    _normalize_rows,
    _retry_after,
    _thinking_config,
    _to_stream_event,
)
from app.services.parsers import Segment


def _api_error(code: int) -> genai_errors.APIError:
    return genai_errors.APIError(code, {"error": {"code": code, "message": "boom", "status": "UNAVAILABLE"}})


def _quota_error(delay: str) -> genai_errors.APIError:
    details = [{"@type": "type.googleapis.com/google.rpc.RetryInfo", "retryDelay": delay}]
    return genai_errors.APIError(
        429, {"error": {"code": 429, "message": "quota", "status": "RESOURCE_EXHAUSTED", "details": details}}
    )


def test_router_orders_models_and_honours_cooldowns() -> None:
    now = [0.0]
    router = ModelRouter(["lite", "main", "other"], clock=lambda: now[0])
    assert router.candidates("main") == ["main", "lite", "other"]
    router.cool_down("main", 60)
    router.cool_down("lite", 60, "+search")
    assert router.candidates("main") == ["lite", "other"]
    assert router.candidates("main", "+search") == ["main", "other"]  # search quota tracked separately
    now[0] = 61
    assert router.candidates("main") == ["main", "lite", "other"]


def test_router_never_returns_empty() -> None:
    router = ModelRouter([])
    router.cool_down("main", 60)
    assert router.candidates("main") == ["main"]


def test_retry_delay_is_parsed_and_bounded() -> None:
    assert _retry_after(_quota_error("42s")) == 42
    assert _retry_after(_quota_error("0.5s")) == 5
    assert _retry_after(_api_error(429)) == 60


class TestGeminiHelpers:
    def test_thinking_config_per_model_family(self) -> None:
        assert _thinking_config("gemini-2.5-flash", 512).thinking_budget == 512
        assert _thinking_config("gemini-3.5-flash", 0).thinking_level == types.ThinkingLevel.MINIMAL
        assert _thinking_config("gemini-3.5-flash", 4096).thinking_level == types.ThinkingLevel.MEDIUM

    def test_lru_cache_evicts_oldest(self) -> None:
        cache = _LRUCache(max_items=2)
        cache.put("a", np.ones(2))
        cache.put("b", np.ones(2))
        cache.get("a")  # refresh "a"
        cache.put("c", np.ones(2))
        assert cache.get("b") is None
        assert cache.get("a") is not None

    def test_rows_are_unit_length(self) -> None:
        rows = _normalize_rows(np.array([[3.0, 4.0], [0.0, 0.0]]))
        assert np.allclose(np.linalg.norm(rows[0]), 1.0)
        assert not np.any(np.isnan(rows))

    def test_stream_event_extracts_text_and_web_grounding(self) -> None:
        chunk = types.GenerateContentResponse(
            candidates=[
                types.Candidate(
                    content=types.Content(
                        role="model",
                        parts=[types.Part(text="thinking...", thought=True), types.Part(text="Answer")],
                    ),
                    grounding_metadata=types.GroundingMetadata(
                        web_search_queries=["deposit law"],
                        grounding_chunks=[
                            types.GroundingChunk(web=types.GroundingChunkWeb(uri="https://x.gov", title="x.gov"))
                        ],
                    ),
                )
            ]
        )
        event = _to_stream_event(chunk)
        assert event.text == "Answer"  # thoughts are never shown
        assert event.web_sources == [("x.gov", "https://x.gov")]
        assert event.search_queries == ["deposit law"]

    def test_missing_key_is_reported(self) -> None:
        with pytest.raises(LLMError):
            GeminiClient(Settings(_env_file=None, gemini_api_key=""))

    async def test_overload_falls_back_to_lighter_model(self) -> None:
        client = GeminiClient(Settings(_env_file=None, gemini_api_key="k", gemini_fallback_models="lite"))
        tried: list[str] = []

        async def call(model: str) -> str:
            tried.append(model)
            if model != "lite":
                raise _api_error(503)
            return "ok"

        assert await client._with_fallback("main", call) == "ok"
        assert tried == ["main", "lite"]

    async def test_quota_error_puts_model_on_cooldown(self) -> None:
        client = GeminiClient(Settings(_env_file=None, gemini_api_key="k", gemini_fallback_models="lite"))
        tried: list[str] = []

        async def call(model: str) -> str:
            tried.append(model)
            if model == "main":
                raise _quota_error("30s")
            return "ok"

        assert await client._with_fallback("main", call) == "ok"
        assert await client._with_fallback("main", call) == "ok"
        assert tried == ["main", "lite", "lite"]  # second request skips the exhausted model

    async def test_chat_degrades_to_no_search_when_search_quota_is_gone(self, monkeypatch: pytest.MonkeyPatch) -> None:
        client = GeminiClient(Settings(_env_file=None, gemini_api_key="k", gemini_fallback_models=""))
        attempts: list[bool] = []

        async def fake_stream(*, model: str, contents: object, config: types.GenerateContentConfig) -> object:
            attempts.append(bool(config.tools))
            if config.tools:
                raise _quota_error("10s")

            async def chunks():
                yield types.GenerateContentResponse(
                    candidates=[types.Candidate(content=types.Content(role="model", parts=[types.Part(text="Hi")]))]
                )

            return chunks()

        monkeypatch.setattr(client._client.aio.models, "generate_content_stream", fake_stream)
        events = [
            e async for e in client.stream_chat(system="s", messages=[ChatMessage("user", "q")], web_search=True,
                                                thinking_budget=0)
        ]  # fmt: skip
        assert [e.text for e in events] == ["Hi"]
        assert attempts == [True, False]

    async def test_network_timeout_moves_to_the_next_model(self) -> None:
        client = GeminiClient(Settings(_env_file=None, gemini_api_key="k", gemini_fallback_models="lite"))

        async def call(model: str) -> str:
            if model == "main":
                raise httpx.ReadTimeout("model hung")
            return "ok"

        assert await client._with_fallback("main", call) == "ok"

    async def test_client_errors_are_not_retried_and_are_friendly(self) -> None:
        client = GeminiClient(Settings(_env_file=None, gemini_api_key="k"))
        tried: list[str] = []

        async def call(model: str) -> str:
            tried.append(model)
            raise _api_error(400)

        with pytest.raises(LLMError, match="could not process"):
            await client._with_fallback("main", call)
        assert tried == ["main"]


def test_markdown_to_docx_structure() -> None:
    data = markdown_to_docx("Brief", "# Title\n\nPlain **bold** text\n\n- item\n\n| A | B |\n|---|---|\n| 1 | 2 |")
    document = Document(io.BytesIO(data))
    texts = [p.text for p in document.paragraphs]
    assert texts[0] == "Brief"
    assert "Plain bold text" in texts
    assert document.tables[0].cell(1, 1).text == "2"
    assert "bold" in [r.text for p in document.paragraphs for r in p.runs if r.bold]


def test_marked_text_truncates_with_flag() -> None:
    doc = StoredDocument(
        id="d", name="n", kind="pdf", size_bytes=1, sha256="h",
        segments=[Segment("a" * 500, "Page 1"), Segment("b" * 500, "Page 2")], chunks=[], vectors=None,
    )  # fmt: skip
    full, cut = doc.marked_text(10_000)
    assert not cut
    assert full.startswith("[Page 1]\n")
    assert "[Page 2]" in full
    partial, cut = doc.marked_text(800)
    assert cut
    assert "[Page 2]" in partial
    assert len(partial) <= 810
